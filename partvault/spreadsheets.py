"""Collection workbook serialization and side-effect-free import planning."""

import csv
from contextlib import nullcontext
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .models import (
    Category,
    Collection,
    Item,
    Manufacturer,
    Status,
    Tag,
)

COLUMNS = (
    "asset_tag",
    "name",
    "category",
    "location",
    "manufacturer",
    "model",
    "revision",
    "serial",
    "status",
    "parent_asset_tag",
    "tags",
    "notes",
    "manufacture_date",
    "release_date",
    "acquired_on",
    "last_tested_on",
)
DATES = {"manufacture_date", "release_date", "acquired_on", "last_tested_on"}
RELATIONS = {
    "category": Category,
    "manufacturer": Manufacturer,
    "status": Status,
    "tags": Tag,
}
SCALARS = tuple(
    c for c in COLUMNS if c not in {*RELATIONS, "asset_tag", "parent_asset_tag"}
)
TEMP_ID = re.compile(r"new:[A-Za-z0-9_-]{1,60}\Z")
ASSET_ID = re.compile(r"[0-9A-Z]{6}\Z")
QUALIFIED = re.compile(r"(.*) \[id:(\d+)\]\Z", re.DOTALL)


class SpreadsheetError(ValueError):
    pass


def limits():
    return (
        getattr(settings, "SPREADSHEET_MAX_UPLOAD_BYTES", 10 * 1024 * 1024),
        getattr(settings, "SPREADSHEET_MAX_ROWS", 5000),
    )


def catalog_for(collection):
    return {
        field: list(
            model.objects.filter(
                Q(user=collection.owner) | Q(user__isnull=True)
            ).order_by("pk")
        )
        for field, model in RELATIONS.items()
    }


def label(obj, catalog):
    if (
        QUALIFIED.fullmatch(obj.name)
        or sum(entry.name == obj.name for entry in catalog) > 1
    ):
        return f"{obj.name} [id:{obj.pk}]"
    return obj.name


def tags_text(values):
    stream = StringIO()
    csv.writer(stream, lineterminator="\n").writerow(values)
    return stream.getvalue()[:-1]


def item_values(item, catalog):
    values = {field: getattr(item, field) for field in SCALARS}
    values.update(
        asset_tag=item.asset_tag or "",
        parent_asset_tag=item.parent_item.asset_tag if item.parent_item else "",
    )
    for field in ("category", "manufacturer", "status"):
        obj = getattr(item, field)
        values[field] = label(obj, catalog[field]) if obj else ""
    values["tags"] = tags_text(
        sorted(label(tag, catalog["tags"]) for tag in item.tags.all())
    )
    return values


def collection_items(collection):
    return (
        Item.objects.filter(collection=collection)
        .select_related("category", "manufacturer", "status", "parent_item")
        .prefetch_related("tags")
        .order_by("asset_tag")
    )


def workbook_bytes(collection, *, template=False):
    catalog = catalog_for(collection)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Items"
    sheet.append(COLUMNS)
    if not template:
        for item in collection_items(collection):
            values = item_values(item, catalog)
            sheet.append([values[column] for column in COLUMNS])
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:P{max(sheet.max_row, 1)}"
    for index, column in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            45 if column == "notes" else 23
        )
        for cells in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            cell = cells[0]
            cell.number_format = "yyyy-mm-dd" if column in DATES else "@"
        # Also format blank template cells so Excel preserves leading zeroes.
        if template:
            for row in range(2, 102):
                sheet.cell(row, index).number_format = (
                    "yyyy-mm-dd" if column in DATES else "@"
                )
    instructions = workbook.create_sheet("Instructions")
    for line in [
        "Part Vault collection spreadsheet",
        f"Collection: {collection.name}",
        "Edit the Items sheet. Keep every header; column order may change. Do not add columns.",
        "Blank asset_tag creates a new item. An existing six-character asset tag updates an item in this collection.",
        "Use new:computer for a new item that other rows reference. Temporary IDs use letters, digits, underscores, or hyphens (1–60 characters).",
        "parent_asset_tag accepts an existing tag or a new:identifier from any row. Blank removes the parent.",
        "Blank cells clear values. A blank name is generated from manufacturer and model, as in the item editor.",
        'Tags are comma separated. Quote a tag containing commas, for example: vintage,"parts, spare".',
        "Reference lists available values. Use qualified labels when names are ambiguous. Missing names become personal values when imported.",
        "Dates must be Excel dates or YYYY-MM-DD. Keep asset tags and serial numbers as text, including leading zeroes.",
        "Formula cells are not accepted. Paste values instead. Entirely blank rows are ignored.",
        "Removing rows does not delete items. Photos, documents, and links are not edited.",
        "Choose Preview changes to inspect without saving, or Import now to validate and save. After previewing, upload the file again to import.",
        "Each import creates items again for blank asset tags and new: identifiers. Export after importing to obtain permanent tags for future edits.",
        f"Upload limit: {limits()[0] // (1024 * 1024)} MiB and {limits()[1]} nonempty rows.",
    ]:
        instructions.append([line])
    instructions.column_dimensions["A"].width = 125
    reference = workbook.create_sheet("Reference")
    reference.append(["field", "accepted value", "scope"])
    for field, entries in catalog.items():
        for obj in entries:
            reference.append(
                [
                    field,
                    label(obj, entries),
                    "Shared" if obj.user_id is None else "Personal",
                ]
            )
    for name in ("A", "B", "C"):
        reference.column_dimensions[name].width = 40
    # openpyxl treats strings starting with '=' as formulas unless explicitly typed.
    for worksheet in workbook:
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.row == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="284B63")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_workbook(upload):
    max_bytes, max_rows = limits()
    if not upload.name.lower().endswith(".xlsx"):
        raise SpreadsheetError("Upload an Excel .xlsx file.")
    if upload.size > max_bytes:
        raise SpreadsheetError(
            f"Workbook exceeds the {max_bytes // (1024 * 1024)} MiB upload limit."
        )
    data = upload.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise SpreadsheetError("Workbook exceeds the upload limit.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise SpreadsheetError(
                    "Workbook expands beyond the 100 MiB processing limit."
                )
        workbook = load_workbook(
            BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if "Items" not in workbook.sheetnames:
                raise SpreadsheetError("Workbook must contain an Items sheet.")
            sheet = workbook["Items"]
            # Do not trust dimensions embedded by third-party spreadsheet writers.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            header = next(iterator, ())
            headers = [
                str(cell.value).strip() if cell.value is not None else ""
                for cell in header
            ]
            while headers and not headers[-1]:
                headers.pop()
            if len(headers) != len(COLUMNS) or set(headers) != set(COLUMNS):
                raise SpreadsheetError(
                    "Headers must contain each template column exactly once, with no additional columns."
                )
            rows = []
            for row_number, cells in enumerate(iterator, 2):
                if row_number > max_rows * 10 + 1:
                    raise SpreadsheetError(
                        "Too many worksheet rows; remove unused rows and upload again."
                    )
                if all(cell.value is None or cell.value == "" for cell in cells):
                    continue
                if len(rows) >= max_rows:
                    raise SpreadsheetError(
                        f"Workbook exceeds the {max_rows} nonempty row limit."
                    )
                values = {}
                errors = []
                if any(cell.value is not None for cell in cells[len(headers) :]):
                    errors.append("Unexpected values outside the template columns.")
                for index, field in enumerate(headers):
                    cell = cells[index] if index < len(cells) else None
                    value = cell.value if cell else None
                    if cell and cell.data_type in ("f", "e"):
                        errors.append(
                            f"{field}: formulas and Excel errors are not accepted."
                        )
                        value = None
                    if isinstance(value, (date, datetime)):
                        value = (
                            value.date().isoformat()
                            if isinstance(value, datetime)
                            else value.isoformat()
                        )
                    elif value is not None:
                        if field in (
                            "asset_tag",
                            "parent_asset_tag",
                        ) and not isinstance(value, str):
                            errors.append(
                                f"{field}: format identifiers as text, preserving all six characters."
                            )
                        value = str(value)
                    values[field] = (
                        (value or "").strip() if field != "notes" else (value or "")
                    )
                rows.append({"number": row_number, "values": values, "errors": errors})
            if not rows:
                raise SpreadsheetError("The Items sheet contains no item rows.")
            return rows
        finally:
            workbook.close()
    except SpreadsheetError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        ValueError,
        TypeError,
        OSError,
        SyntaxError,
    ) as exc:
        raise SpreadsheetError(
            "Unable to read this workbook. Save it as a valid .xlsx file and try again."
        ) from exc


def identifier(value):
    return value if value.startswith("new:") else value.upper()


def resolve_value(value, field, catalog):
    if not value:
        return None
    model = RELATIONS[field]
    qualified = QUALIFIED.fullmatch(value)
    if qualified:
        name, pk = qualified.groups()
        matches = [
            obj for obj in catalog[field] if obj.pk == int(pk) and obj.name == name
        ]
        if not matches:
            raise SpreadsheetError(
                f"{field}: reference label is unavailable; use the Reference sheet."
            )
    else:
        matches = [obj for obj in catalog[field] if obj.name == value]
    if len(matches) > 1:
        raise SpreadsheetError(
            f"{field}: '{value}' is ambiguous; use a qualified label from the Reference sheet."
        )
    if matches:
        return {"id": matches[0].pk, "name": matches[0].name}
    try:
        model._meta.get_field("name").clean(value, None)
    except ValidationError as exc:
        raise SpreadsheetError(f"{field}: {'; '.join(exc.messages)}") from exc
    return {"id": None, "name": value}


def display(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value or "")


def normalize_line_endings(value):
    return value.replace("\r\n", "\n").replace("\r", "\n")


def build_plan(collection, rows):
    catalog = catalog_for(collection)
    items = list(collection_items(collection))
    by_tag = {item.asset_tag: item for item in items if item.asset_tag}
    parents = {
        item.asset_tag: item.parent_item.asset_tag if item.parent_item else ""
        for item in items
        if item.asset_tag
    }
    planned = []
    seen = {}
    new_values = {field: set() for field in RELATIONS}
    for row in rows:
        values = row["values"]
        asset = identifier(values["asset_tag"])
        parent = identifier(values["parent_asset_tag"])
        errors = list(row["errors"])
        is_new = not asset or asset.startswith("new:")
        existing = by_tag.get(asset) if not is_new else None
        if asset:
            if asset in seen:
                errors.append(
                    f"asset_tag: duplicate identifier (also row {seen[asset]})."
                )
            seen[asset] = row["number"]
            if is_new and not TEMP_ID.fullmatch(asset):
                errors.append(
                    "asset_tag: use new: followed by 1–60 letters, digits, underscores, or hyphens."
                )
            elif not is_new and (not ASSET_ID.fullmatch(asset) or existing is None):
                errors.append("asset_tag: no matching item in this collection.")
        key = asset or f"row:{row['number']}"
        relations = {}
        for field in RELATIONS:
            try:
                if field == "tags":
                    names = (
                        next(
                            csv.reader(
                                [values[field]], skipinitialspace=True, strict=True
                            )
                        )
                        if values[field]
                        else []
                    )
                    relations[field] = [
                        resolve_value(name, field, catalog)
                        for name in dict.fromkeys(n.strip() for n in names)
                        if name
                    ]
                else:
                    relations[field] = resolve_value(values[field], field, catalog)
            except (SpreadsheetError, csv.Error) as exc:
                errors.append(str(exc))
                relations[field] = [] if field == "tags" else None
        scalars = {}
        for field in SCALARS:
            value = values[field]
            try:
                if field in DATES:
                    if value and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                        raise ValidationError("Use an Excel date or YYYY-MM-DD.")
                    value = date.fromisoformat(value) if value else None
                value = Item._meta.get_field(field).clean(value, None)
                scalars[field] = display(value)
            except (ValueError, ValidationError) as exc:
                messages = (
                    exc.messages
                    if isinstance(exc, ValidationError)
                    else ["Invalid date."]
                )
                errors.append(f"{field}: {'; '.join(messages)}")
                scalars[field] = ""
        if not scalars["name"].strip():
            manufacturer = relations.get("manufacturer")
            scalars["name"] = " ".join(
                filter(
                    None,
                    [manufacturer["name"] if manufacturer else "", scalars["model"]],
                )
            )
            if len(scalars["name"]) > 200:
                errors.append(
                    "name: generated name exceeds 200 characters; provide a shorter name."
                )
        changes = []
        previous = item_values(existing, catalog) if existing else {}
        if existing and normalize_line_endings(
            previous["notes"]
        ) == normalize_line_endings(scalars["notes"]):
            # Excel normalizes newlines. Preserve the original notes even when
            # another field on this item is being updated.
            scalars["notes"] = previous["notes"]
        for field in SCALARS:
            if not existing or display(previous[field]) != scalars[field]:
                changes.append(
                    {
                        "field": field,
                        "before": display(previous.get(field)),
                        "after": scalars[field],
                    }
                )
        for field, selected in relations.items():
            entries = selected if field == "tags" else ([selected] if selected else [])
            for entry in entries:
                if entry["id"] is None:
                    new_values[field].add(entry["name"])
            old_ids = (
                set(tag.pk for tag in existing.tags.all())
                if existing and field == "tags"
                else {getattr(existing, f"{field}_id", None)}
                if existing
                else set()
            )
            new_ids = (
                {entry["id"] for entry in entries}
                if field == "tags"
                else {selected["id"] if selected else None}
            )
            if (
                not existing
                or old_ids != new_ids
                or any(entry["id"] is None for entry in entries)
            ):
                changes.append(
                    {
                        "field": field,
                        "before": display(previous.get(field)),
                        "after": values[field],
                    }
                )
        old_parent = previous.get("parent_asset_tag", "")
        if not existing or parent != old_parent:
            changes.append(
                {"field": "parent_asset_tag", "before": old_parent, "after": parent}
            )
        planned.append(
            {
                "number": row["number"],
                "key": key,
                "asset_tag": asset,
                "item_id": existing.pk if existing else None,
                "action": "create" if is_new else "update" if changes else "unchanged",
                "scalars": scalars,
                "relations": relations,
                "parent": parent,
                "changes": changes,
                "errors": errors,
            }
        )
        parents[key] = parent
    available = set(by_tag) | {
        row["key"] for row in planned if row["action"] == "create" and row["asset_tag"]
    }
    for row in planned:
        if row["parent"] and row["parent"] not in available:
            row["errors"].append(
                "parent_asset_tag: no matching existing item or temporary identifier."
            )
        visited = set()
        current = row["key"]
        while current and current in parents:
            if current in visited:
                row["errors"].append(
                    "parent_asset_tag: relationship creates or reaches a parent cycle."
                )
                break
            visited.add(current)
            current = parents[current]
    counts = {
        action: sum(row["action"] == action for row in planned)
        for action in ("create", "update", "unchanged")
    }
    counts["errors"] = sum(bool(row["errors"]) for row in planned)
    return {
        "rows": planned,
        "counts": counts,
        "new_values": {
            field: sorted(names) for field, names in new_values.items() if names
        },
    }


def process_workbook(collection, user, upload, *, commit=False):
    """Validate each upload afresh; optionally apply its plan atomically."""
    rows = read_workbook(upload)
    with transaction.atomic() if commit else nullcontext():
        collections = Collection.objects.filter(pk=collection.pk, owner=user)
        if commit:
            collections = collections.select_for_update()
        collection = collections.get()
        if commit:
            list(
                Item.objects.select_for_update()
                .filter(collection=collection)
                .values_list("pk", flat=True)
            )
        plan = build_plan(collection, rows)
        if commit and not plan["counts"]["errors"]:
            _apply_plan(collection, user, plan)
        return plan


def _apply_plan(collection, user, plan):
    """Apply a validated plan inside the caller's transaction."""
    related = {}
    for field, model in RELATIONS.items():
        for name in plan["new_values"].get(field, []):
            obj, _ = model.objects.get_or_create(user=user, name=name)
            related[(field, name)] = obj.pk

    def related_id(field, value):
        return value["id"] or related[(field, value["name"])] if value else None

    all_items = {
        item.asset_tag: item for item in Item.objects.filter(collection=collection)
    }
    for row in plan["rows"]:
        if row["action"] == "unchanged":
            continue
        item = (
            all_items[row["asset_tag"]]
            if row["action"] == "update"
            else Item(collection=collection)
        )
        for field, value in row["scalars"].items():
            setattr(
                item,
                field,
                date.fromisoformat(value)
                if field in DATES and value
                else None
                if field in DATES
                else value,
            )
        for field in ("category", "manufacturer", "status"):
            setattr(item, f"{field}_id", related_id(field, row["relations"][field]))
        item.save()
        item.tags.set([related_id("tags", value) for value in row["relations"]["tags"]])
        all_items[row["key"]] = item
        row["asset_tag"] = item.asset_tag
        row["item_id"] = item.pk
    for row in plan["rows"]:
        if row["action"] != "unchanged":
            item = all_items[row["key"]]
            parent = all_items.get(row["parent"])
            if item.parent_item_id != (parent.pk if parent else None):
                item.parent_item = parent
                item.save(update_fields=["parent_item"])
            for change in row["changes"]:
                if change["field"] == "parent_asset_tag":
                    change["after"] = parent.asset_tag if parent else ""
